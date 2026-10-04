import json
import re
from collections import defaultdict
from pathlib import Path

from .domain import clone, stamp


def ordered(project, rows, grouped=False, character=None):
    if character:
        selected = []
        for row in rows:
            r = clone(row)
            if r["turns"]:
                r["turns"] = [t for t in r["turns"] if t.get("character_id") == character]
                if not r["turns"]:
                    continue
                r["translation"] = "\n".join(t.get("text", "") for t in r["turns"])
                r["character_ids"] = [character]
                r["notes"] = "整段原文含其他人物；以下中文仅为选定人物的分句。 " + r["notes"]
            elif character not in r["character_ids"]:
                continue
            selected.append(r)
        rows = selected
    if not grouped:
        return [("时间顺序", rows)]
    groups = defaultdict(list)
    for r in rows:
        if r.get("turns"):
            for cid in dict.fromkeys(t.get("character_id") for t in r["turns"]):
                part = clone(r)
                part["turns"] = [t for t in r["turns"] if t.get("character_id") == cid]
                part["translation"] = "\n".join(t.get("text", "") for t in part["turns"])
                part["character_ids"] = [cid] if cid else []
                part["notes"] = (
                    "整段原文含其他人物；分句时间未填写时使用整段定位时间。 " + part["notes"]
                )
                groups[project.role(part)].append(part)
        else:
            groups[project.role(r)].append(r)
    return sorted(groups.items(), key=lambda x: (x[0] == "待确认", -len(x[1]), x[0]))


def item(project, r, episodes):
    a = f"{episodes.get(r['episode_id'], '')} [{stamp(r['start_ms'])} – {stamp(r['end_ms'])}] {project.role(r)} / {r['kind']}\n"
    a += f"人物：{'人工确认' if r['speaker_review'] == 'confirmed' else '待复核'}；文字：{'已校对' if r['text_review'] == 'confirmed' else '待校对'}\n"
    a += "原文：" + r["original"] + "\n中文：" + r["translation"] + "\n"
    if r["turns"]:
        names = {c["id"]: c["name"] for c in project.characters()}
        a += (
            "分句归属：\n"
            + "\n".join(
                names.get(t.get("character_id"), "待确认") + "：" + t.get("text", "")
                for t in r["turns"]
            )
            + "\n"
        )
    a += "依据：" + "；".join(r["evidence"]) + "\n"
    if r["flags"]:
        a += "识别备注：" + "；".join(r["flags"]) + "\n"
    if r["notes"]:
        a += "备注：" + r["notes"] + "\n"
    return a


def export_txt(project, rows, path, grouped=False, character=None):
    episodes = {e["id"]: e["title"] for e in project.episodes()}
    content = (
        project.meta("title")
        + " 台词本\n自动识别及人物候选需核对；人工审核与文字校对分别标记。\n\n"
    )
    for name, rs in ordered(project, rows, grouped, character):
        content += (
            name
            + "\n"
            + "=" * 30
            + "\n\n"
            + "\n".join(item(project, r, episodes) for r in rs)
            + "\n"
        )
    Path(path).write_text(content, encoding="utf-8-sig")


def export_docx(project, rows, path, grouped=False, character=None):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.65)
    for name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        for border in list(style._element.findall(".//" + qn("w:pBdr"))):
            border.getparent().remove(border)
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Normal"].paragraph_format.space_after = Pt(4)
    doc.add_paragraph(project.meta("title") + " 台词本", style="Title")
    doc.add_paragraph(
        "本台词本整理视频原声与中文对照，并保留人物归属及审核状态。自动结果仍需核对；人物确认不代表原文或译文已经校对。"
    )
    episodes = {e["id"]: e["title"] for e in project.episodes()}
    for name, rs in ordered(project, rows, grouped, character):
        doc.add_heading(name, level=1)
        table = doc.add_table(rows=1, cols=4)
        table.autofit = False
        widths = [1.25, 1.25, 2.35, 2.35]
        for col, width in zip(table.columns, widths):
            col.width = Inches(width)
        for i, (cell, label) in enumerate(
            zip(table.rows[0].cells, ["剧集与时间", "人物与审核", "原声文字", "中文对照"])
        ):
            cell.text = label
            cell.width = Inches(widths[i])
            sh = OxmlElement("w:shd")
            sh.set(qn("w:fill"), "EAF0F6")
            cell._tc.get_or_add_tcPr().append(sh)
        repeat = OxmlElement("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for r in rs:
            cells = table.add_row().cells
            cells[0].text = (
                episodes.get(r["episode_id"], "")
                + "\n"
                + stamp(r["start_ms"])
                + "\n"
                + stamp(r["end_ms"])
            )
            cells[1].text = (
                project.role(r)
                + "\n"
                + ("人物已确认" if r["speaker_review"] == "confirmed" else "人物待复核")
                + "\n"
                + ("文字已校对" if r["text_review"] == "confirmed" else "文字待校对")
            )
            cells[2].text = r["original"]
            cells[3].text = r["translation"]
            if r["notes"]:
                cells[3].add_paragraph(r["notes"])
            for i, cell in enumerate(cells):
                cell.width = Inches(widths[i])
        for row in table.rows:
            cant_split = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(cant_split)
            for cell in row.cells:
                from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                props = cell._tc.get_or_add_tcPr()
                borders = OxmlElement("w:tcBorders")
                for edge in ["top", "left", "bottom", "right"]:
                    node = OxmlElement("w:" + edge)
                    node.set(qn("w:val"), "single")
                    node.set(qn("w:sz"), "4")
                    node.set(qn("w:color"), "D9D9D9")
                    borders.append(node)
                props.append(borders)
                margins = OxmlElement("w:tcMar")
                for edge in ["top", "left", "bottom", "right"]:
                    node = OxmlElement("w:" + edge)
                    node.set(qn("w:w"), "90")
                    node.set(qn("w:type"), "dxa")
                    margins.append(node)
                props.append(margins)
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1.1
                    for run in p.runs:
                        run.font.size = Pt(9.5)
    doc.save(path)


def export_srt(project, rows, path, language="dual", speaker=True):
    episodes = defaultdict(list)
    for r in rows:
        if r["kind"] not in ["画面文字"]:
            episodes[r["episode_id"]].append(r)
    target = Path(path)
    names = {e["id"]: e["title"] for e in project.episodes()}
    written = []
    for eid, rs in episodes.items():
        if len(episodes) == 1:
            dest = target
        else:
            number = next(e["number"] for e in project.episodes() if e["id"] == eid)
            dest = target.with_name(
                target.stem + f"_{number:02}_" + re.sub(r'[<>:"/\\|?*]', "_", names[eid]) + ".srt"
            )
        blocks = []
        for r in sorted(rs, key=lambda r: r["start_ms"]):
            chinese = (
                r["original"]
                if r.get("language") in ["zh", "chinese"] and not r["translation"]
                else r["translation"]
            )
            text = (
                r["original"]
                if language == "original"
                else chinese
                if language == "zh"
                else "\n".join(s for s in [r["original"], r["translation"]] if s)
            )
            if not text:
                continue
            if speaker:
                text = project.role(r) + "：" + text
            blocks.append(
                f"{len(blocks) + 1}\n{stamp(r['start_ms'], True)} --> {stamp(r['end_ms'], True)}\n{text}\n"
            )
        dest.write_text("\n".join(blocks), encoding="utf-8-sig")
        written.append(dest)
    return written


def payload(project, rows):
    data = clone(rows)
    for r in data:
        r["base_revision"] = r["revision"]
    return dict(
        schema=1,
        project_id=project.meta("id"),
        title=project.meta("title"),
        rows=data,
        characters=project.characters(),
        episodes=project.episodes(),
    )


def export_json(project, rows, path):
    Path(path).write_text(
        json.dumps(payload(project, rows), ensure_ascii=False, indent=2), encoding="utf8"
    )


def export_html(project, rows, path, grouped=False):
    template = (Path(__file__).parent / "assets" / "review.html").read_text(encoding="utf8")
    data = payload(project, rows)
    data["grouped"] = grouped
    text = json.dumps(data, ensure_ascii=False).replace("<", "\\u003c")
    Path(path).write_text(template.replace("__PROJECT__", text), encoding="utf8")
